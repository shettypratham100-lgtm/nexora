"""
text_extractor.py

This module is responsible for extracting plain text from supported
document formats.

Currently Supported:
- PDF
- DOCX
- PPTX

Future Support:
- Images (OCR)
- TXT
- XLSX
"""

from pathlib import Path

import fitz                    # PyMuPDF
from docx import Document
from pptx import Presentation

def extract_text(file_path):
    """
    Extract plain text from a supported document.

    Parameters
    ----------
    file_path : str or Path
        Absolute path of the uploaded file.

    Returns
    -------
    str
        Extracted plain text.

    Raises
    ------
    ValueError
        If the file type is unsupported.
    """

    file_path = Path(file_path)

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    elif extension == ".docx":
        return extract_docx_text(file_path)

    elif extension in [".ppt", ".pptx"]:
        return extract_ppt_text(file_path)

    raise ValueError(f"Unsupported file type: {extension}")

def extract_pdf_text(file_path):
    """
    Extract text from a PDF document.
    """

    document = fitz.open(file_path)

    extracted_text = []

    try:

        for page in document:
            extracted_text.append(page.get_text())

    finally:
        document.close()

    return "\n".join(extracted_text).strip()

def extract_docx_text(file_path):
    """
    Extract text from a Microsoft Word document.
    """

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        paragraphs.append(paragraph.text)

    return "\n".join(paragraphs).strip()

def extract_ppt_text(file_path):
    """
    Extract text from a PowerPoint presentation.
    """

    presentation = Presentation(file_path)

    extracted_text = []

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                extracted_text.append(shape.text)

    return "\n".join(extracted_text).strip()